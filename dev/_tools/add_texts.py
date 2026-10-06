"""Ajoute des textes du mod dans les deux fichiers de langue. usage : python _tools/add_texts.py fichier.json
Le fichier : une liste de [identifiant, japonais, anglais, chinois]. Un identifiant deja present est mis a jour."""
import io
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "package/LangMod/EN/emp_localization.xlsx"
CN = ROOT / "package/LangMod/CN/SourceLocalization.json"


def main():
    texts = json.load(io.open(sys.argv[1], encoding="utf-8"))
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["General"]
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    for key, jp, en, _ in texts:
        if key in rows:
            ws.cell(rows[key], 3).value, ws.cell(rows[key], 4).value = jp, en
        else:
            ws.append([key, None, jp, en])
    wb.save(XLSX)

    raw = io.open(CN, encoding="utf-8").read()
    data = json.loads(raw)
    for key, _, _, cn in texts:
        data[f"LangGeneral.{key}.text"] = cn
    indent = 4 if raw.startswith("{\n    ") else 2
    io.open(CN, "w", encoding="utf-8", newline="\n").write(
        json.dumps(data, ensure_ascii=False, indent=indent) + ("\n" if raw.endswith("\n") else ""))
    print(len(texts), "textes")


if __name__ == "__main__":
    main()
